#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import difflib
import html as htmllib
import json
import os
import random
import re
import shutil
import sys
import time
import unicodedata
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from datetime import datetime, timedelta
from urllib.parse import urlparse, urlunparse, parse_qsl, urlencode, quote

import pandas as pd
import trafilatura
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment

try:                                   # package mới tên là "ddgs"
    from ddgs import DDGS
    HAVE_DDGS = True
except ImportError:
    from duckduckgo_search import DDGS
    HAVE_DDGS = False

# ======================= CẤU HÌNH =======================
DISASTER_FILE = "[8386] DISASTER DATA.xlsx"
DISASTER_SHEET = "(ĐÃ LỌC) THEO QUỐC GIA"
NEWS_TEMPLATE = "[8386] NEWS DATA.xlsx"                  # File mẫu: KHÔNG bị ghi đè
OUTPUT_PATTERN = "[8386] NEWS DATA - {country}.xlsx"     # Mỗi nước 1 file
MERGED_OUTPUT = "[8386] NEWS DATA - MERGED.xlsx"

# ---- KHÔNG GIỚI HẠN số bài: cào hết mọi đài của nước, mọi query ----
RESULTS_PER_QUERY = 30      # số link mỗi query (công cụ tìm kiếm có thể trả ít hơn)
MAX_QUERIES_PER_OUTLET = None   # None = chạy hết mọi biến thể query trên mỗi đài
MAX_PLACES = 3              # số địa danh đầu tiên trong cột Location dùng để tạo query
MIN_WORDS = 100
DAYS_BEFORE, DAYS_AFTER = 7, 45   # cửa sổ ngày hợp lệ quanh thảm họa
FIRST_DATA_ROW = 3

DOWNLOAD_WORKERS = 6        # số bài tải song song
USE_WP_API = True           # Đài dùng WordPress: lấy bài qua API có lọc ngày. Tự phát hiện, không có thì bỏ qua
WP_PER_PAGE = 50
WP_MAX_PAGES = 10           # số trang API tối đa cho mỗi từ khóa
AUTO_ARTICLE_ID = False     # False: để trống cột Article ID (đánh tay). True: tự đánh ID001, VN001...
NAMED_EVENT_STRICT = False  # False: nhận bài nếu nhắc tên bão HOẶC (đúng loại + địa danh). True: bão có tên bắt buộc nhắc tên (luật cũ)
STRICT_PLACE = True         # Bài (nhánh đúng loại thiên tai) phải nhắc ít nhất 1 địa danh trong Location, hoặc tiêu đề nói về loại thiên tai
SLEEP_RANGE = (2, 4)        # nghỉ giữa các lần search (giây), tránh bị chặn
SEARCH_BACKENDS = ["duckduckgo", "bing", "brave", "google"] if HAVE_DDGS else [None]

# ---- Chống trùng/lẫn bài giữa các sự kiện cùng nước, cùng khoảng thời gian ----
AMBIGUOUS_MARGIN = 1.5      # chênh điểm giữa sự kiện nhất và nhì < ngưỡng này -> ghi vào ambiguous_<NƯỚC>.csv để kiểm tay
OTHER_NAMED_PENALTY = 8.0   # bài nhắc tên bão của sự kiện khác -> các sự kiện không nhắc tên đó bị trừ điểm

COUNTRY_INFO = {
    "INDONESIA":   {"code": "IDN", "prefix": "ID", "name": "Indonesia"},
    "PHILIPPINES": {"code": "PHL", "prefix": "PH", "name": "Philippines"},
    "THAILAND":    {"code": "THA", "prefix": "TH", "name": "Thailand"},
    "MALAYSIA":    {"code": "MYS", "prefix": "MY", "name": "Malaysia"},
    "VIETNAM":     {"code": "VNM", "prefix": "VN", "name": "Vietnam"},
}

# Danh sách đài theo thứ tự cào của từng nước (đài chính trước). Cào HẾT tất cả.
OUTLETS = {
    "INDONESIA": [
        ("The Jakarta Post", "thejakartapost.com"),
        ("ANTARA News - English", "antaranews.com"),
        ("Tempo English", "en.tempo.co"),
        ("Jakarta Globe", "jakartaglobe.id"),
        ("RRI - Voice of Indonesia", "rri.co.id/voice-of-indonesia"),
        ("Indonesia Expat", "indonesiaexpat.id"),
        ("The Bali Times", "thebalitimes.com"),
        ("Inside Indonesia", "insideindonesia.org"),
        ("The Conversation Indonesia", "theconversation.com"),
        ("Kompas (English)", "kompas.id"),
    ],
    "PHILIPPINES": [
        ("Philippine Daily Inquirer", "inquirer.net"),
        ("GMA News Online", "gmanetwork.com"),
        ("ABS-CBN News", "abs-cbn.com"),
        ("Philstar.com", "philstar.com"),
        ("Rappler", "rappler.com"),
        ("Manila Bulletin", "mb.com.ph"),
        ("BusinessWorld", "bworldonline.com"),
        ("The Manila Times", "manilatimes.net"),
        ("Manila Standard", "manilastandard.net"),
    ],
    "THAILAND": [
        ("Bangkok Post", "bangkokpost.com"),
        ("The Nation Thailand", "nationthailand.com"),
        ("Thai PBS World", "thaipbsworld.com"),
        ("Khaosod English", "khaosodenglish.com"),
        ("Pattaya Mail", "pattayamail.com"),
        ("MCOT", "mcot.net"),
        ("NNT - National News Bureau of Thailand", "thainationalnews.com"),
        ("The Thaiger", "thethaiger.com"),
        ("Thai Enquirer", "thaienquirer.com"),
    ],
    "MALAYSIA": [
        ("Bernama", "bernama.com"),
        ("The Star", "thestar.com.my"),
        ("New Straits Times", "nst.com.my"),
        ("Malay Mail", "malaymail.com"),
        ("Free Malaysia Today (FMT)", "freemalaysiatoday.com"),
        ("The Borneo Post", "theborneopost.com"),
        ("The Edge Malaysia", "theedgemalaysia.com"),
        ("The Sun Daily", "thesun.my"),
        ("Daily Express", "dailyexpress.com.my"),
    ],
    "VIETNAM": [
        ("VietnamPlus", "en.vietnamplus.vn"),
        ("Vietnam News", "vietnamnews.vn"),
        ("VnExpress International", "e.vnexpress.net"),
        ("Tuoi Tre News", "tuoitrenews.vn"),
        ("VietNamNet Global", "vietnamnet.vn/en"),
        ("VOV English", "english.vov.vn"),
        ("Nhan Dan English", "en.nhandan.vn"),
        ("Vietnam Investment Review (VIR)", "vir.com.vn"),
        ("Saigon Times English", "english.thesaigontimes.vn"),
    ],
}

# Các domain chỉ nhận bài TIẾNG ANH (bài tiếng Indonesia bị loại sau khi tải, dựa vào nội dung bài).
ENGLISH_ONLY_SITES = {"theconversation.com", "kompas.id"}
MIN_EN_RATIO = 0.15         # tỉ lệ từ dừng tiếng Anh tối thiểu trong 400 từ đầu của bài

COLS = ["OUTLET", "Disaster Name", "Article ID", "Tittle", "Full text", "Date", "URL"]

EXCLUDED_DOMAINS = ["login.yahoo", "facebook.com", "youtube.com", "twitter.com",
                    "instagram.com", "accounts.google"]
# Trang không phải bài báo (thẻ, chuyên mục, tìm kiếm, phân trang...)
JUNK_URL_PATTERNS = re.compile(
    r"/(tag|tags|topic|topics|category|categories|search|author|page|video|photo)s?(/|\?|$)", re.I)


# ======================= DANH SÁCH ĐÀI =======================
def find_header_row(ws):
    """Dòng tiêu đề = dòng có ô cột A là 'OUTLET' (mặc định dòng 2). Dữ liệu bắt đầu ngay dưới nó."""
    for r in range(1, 11):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().upper() == "OUTLET":
            return r
    return FIRST_DATA_ROW - 1


def get_outlets(country):
    """[(tên đài, site)] theo thứ tự trong OUTLETS, bỏ đài trùng domain."""
    outlets, used = [], set()
    for name, site in OUTLETS[country]:
        if site in used:
            continue
        used.add(site)
        outlets.append((name, site))
    return outlets


# ======================= TÌM LINK BẰNG DUCKDUCKGO =======================
def url_matches_site(url, site):
    """Link phải thật sự thuộc đài đang cào (site: của công cụ tìm kiếm đôi khi trả link lạ)."""
    p = urlparse(url)
    host = p.netloc.lower().split(":")[0]
    base, _, path = site.partition("/")
    if not (host == base or host.endswith("." + base)):
        return False
    return p.path.startswith("/" + path) if path else True


def norm_url(u):
    """Chuẩn hóa để so trùng: bỏ #fragment, tham số utm_*, dấu / cuối."""
    p = urlparse(u.strip())
    q = [(k, v) for k, v in parse_qsl(p.query) if not k.lower().startswith("utm_")]
    return urlunparse((p.scheme, p.netloc.lower(), p.path.rstrip("/"), "", urlencode(q), ""))


def _good_urls(results, site):
    urls = []
    for r in results or []:
        u = r.get("href", "")
        if (u.startswith("http")
                and url_matches_site(u, site)
                and not any(bad in u.lower() for bad in EXCLUDED_DOMAINS)
                and not JUNK_URL_PATTERNS.search(u)
                and u not in urls):
            urls.append(u)
    return urls


def search_article_urls(query, site, retries=2):
    """Thử lần lượt từng công cụ tìm kiếm.
    Trả về [] nếu query rỗng thật, None nếu tất cả công cụ đều lỗi kết nối."""
    for attempt in range(retries):
        errors, got_empty = [], False
        for backend in SEARCH_BACKENDS:
            try:
                kw = {"backend": backend} if backend else {}
                results = DDGS().text(query, max_results=RESULTS_PER_QUERY, **kw)
                urls = _good_urls(results, site)
                if urls:
                    time.sleep(random.uniform(*SLEEP_RANGE))
                    return urls
                got_empty = True
            except Exception as e:
                if "no results" in str(e).lower():
                    got_empty = True
                else:
                    errors.append(f"{backend or 'default'}: {str(e)[:80]}")
                    print(f"   ⚠️ Công cụ '{backend or 'default'}' lỗi -> thử công cụ khác...")
            time.sleep(random.uniform(0.5, 1.5))
        if got_empty:
            print("      (không có kết quả)")
            return []
        wait = 20 * (attempt + 1)
        print(f"   ⚠️ Tất cả công cụ tìm kiếm đều lỗi ({errors[-1]}). Chờ {wait}s rồi thử lại "
              f"({attempt + 1}/{retries}). Kiểm tra mạng/VPN nếu lỗi kéo dài.")
        time.sleep(wait)
    return None


# ======================= CÀO NỘI DUNG =======================
def clean_text(s):
    if not s:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(s)).strip()[:32000]   # giới hạn 32767 ký tự/ô của Excel


BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/124.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Accept-Encoding": "identity",
}


def download_html(url):
    html = trafilatura.fetch_url(url)
    if html:
        return html
    try:                                              # nhiều báo chặn bot -> thử lại với header giống trình duyệt
        req = urllib.request.Request(url, headers=BROWSER_HEADERS)
        with urllib.request.urlopen(req, timeout=25) as r:
            raw = r.read()
            charset = r.headers.get_content_charset() or "utf-8"
        return raw.decode(charset, errors="replace")
    except Exception:
        return None


def extract_article_content(url):
    """Trả về (bài, lý do_lỗi). lý do: 'tải lỗi' | 'quá ngắn' | 'lỗi đọc'."""
    try:
        downloaded = download_html(url)
        if not downloaded:
            return None, "tải lỗi"
        meta = trafilatura.extract_metadata(downloaded)
        text = trafilatura.extract(downloaded, include_comments=False, include_tables=False)
        if not text or len(text.split()) < MIN_WORDS:
            return None, "quá ngắn"
        return {
            "title": clean_text(meta.title) if meta and meta.title else "No Title",
            "full_text": clean_text(text),
            "date": str(meta.date)[:10] if meta and meta.date else "",
        }, ""
    except Exception:
        return None, "lỗi đọc"


# ---------- Đài chạy WordPress: hỏi thẳng API (lọc theo ngày, có sẵn nội dung) ----------
WP_CACHE = {}


def _http_json(url, timeout=25):
    req = urllib.request.Request(url, headers={**BROWSER_HEADERS, "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8", errors="replace"))


def wp_available(site):
    if "/" in site:                                    # site có path (vd vietnamnet.vn/en) -> không dùng
        return False
    if site in WP_CACHE:
        return WP_CACHE[site]
    try:
        ok = isinstance(_http_json(f"https://{site}/wp-json/wp/v2/posts?per_page=1&_fields=id"), list)
    except Exception:
        ok = False
    WP_CACHE[site] = ok
    return ok


def _strip_tags(h):
    return re.sub(r"\s+", " ", htmllib.unescape(re.sub(r"<[^>]+>", " ", h or ""))).strip()


def wp_search(site, term, window):
    """Tìm bài theo từ khóa trong khoảng ngày của thảm họa. Trả về [{url,title,html,date}]."""
    q = f"https://{site}/wp-json/wp/v2/posts?search={quote(term)}&per_page={WP_PER_PAGE}&_fields=link,date,title,content"
    if window:
        q += f"&after={window[0]:%Y-%m-%d}T00:00:00&before={window[1]:%Y-%m-%d}T23:59:59"
    out = []
    for page in range(1, WP_MAX_PAGES + 1):
        try:
            data = _http_json(q + f"&page={page}")
        except Exception:
            break
        if not isinstance(data, list):
            break
        for d in data:
            try:
                out.append({"url": d["link"], "title": _strip_tags(d["title"]["rendered"]),
                            "html": d["content"]["rendered"], "date": str(d.get("date", ""))[:10]})
            except Exception:
                continue
        if len(data) < WP_PER_PAGE:
            break
        time.sleep(1)
    return out


def wp_to_content(item):
    """Bài WordPress -> (content, lý do lỗi) giống extract_article_content."""
    text = None
    try:
        text = trafilatura.extract(item["html"], include_comments=False, include_tables=False)
    except Exception:
        pass
    text = text or _strip_tags(item["html"])
    if not text or len(text.split()) < MIN_WORDS:
        return None, "quá ngắn"
    return {"title": clean_text(item["title"]) or "No Title", "full_text": clean_text(text),
            "date": item["date"]}, ""


# ======================= TIỆN ÍCH DỮ LIỆU =======================
def _num(x):
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return None
    try:
        return int(float(str(x).strip()))
    except (ValueError, TypeError):
        return None


def _cell(row, *names):
    """Lấy giá trị ô (thử nhiều tên cột), trả '' nếu trống/NaN."""
    for n in names:
        v = row.get(n)
        if v is None or (not isinstance(v, str) and pd.isna(v)):
            continue
        s = str(v).strip()
        if s and s.lower() != "nan":
            return s
    return ""


def _fold(s):
    """Bỏ dấu + lower để so khớp (Quảng Bình == quang binh)."""
    s = unicodedata.normalize("NFD", str(s).replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


_EN_STOP = {"the", "and", "of", "to", "in", "is", "that", "for", "with", "on", "as", "was", "are", "by",
            "at", "from", "it", "this", "be", "have", "has", "said", "were", "which", "an", "their"}
_ID_STOP = {"yang", "dan", "di", "untuk", "dengan", "dari", "pada", "ini", "itu", "adalah", "dalam",
            "akan", "tidak", "juga", "oleh", "karena", "ke", "sebagai", "telah", "para", "atau",
            "saat", "mereka", "ada", "sebuah", "kepada", "bahwa"}


def is_english(text, min_ratio=0.15, sample=400):
    """Đoán bài có phải tiếng Anh không: đếm từ dừng tiếng Anh vs tiếng Indonesia trong `sample` từ đầu.
    Không cần thư viện ngoài. Bài tiếng Indonesia có ít từ dừng tiếng Anh nên bị loại."""
    words = re.findall(r"[a-z']+", _fold(text))[:sample]
    if not words:
        return False
    en = sum(1 for w in words if w in _EN_STOP)
    idn = sum(1 for w in words if w in _ID_STOP)
    return en / len(words) >= min_ratio and en > 2 * idn


def disaster_core(row):
    """(ngày bắt đầu, ngày kết thúc) thật của thảm họa, không cộng đệm."""
    try:
        sy, sm = _num(row.get("Start Year")), _num(row.get("Start Month"))
        if not sy or not sm or not (1 <= sm <= 12):
            return None
        sd = max(1, min(_num(row.get("Start Day")) or 1, 28))
        ey = _num(row.get("End Year")) or sy
        em = max(1, min(_num(row.get("End Month")) or sm, 12))
        ed = max(1, min(_num(row.get("End Day")) or 28, 28))
        return datetime(sy, sm, sd), datetime(ey, em, ed)
    except Exception:
        return None


def disaster_window(row):
    core = disaster_core(row)
    if core is None:
        return None
    return core[0] - timedelta(days=DAYS_BEFORE), core[1] + timedelta(days=DAYS_AFTER)


def date_ok(date_str, window):
    if window is None or not date_str:
        return True  # Không có ngày -> giữ lại để kiểm tay
    try:
        d = datetime.strptime(str(date_str)[:10], "%Y-%m-%d")
    except ValueError:
        return True
    return window[0] <= d <= window[1]


PLACE_NOISE = re.compile(r"\b(provinces?|districts?|cities|city|towns?|regency|regencies|"
                         r"prefecture|region|areas?|nationwide)\b", re.I)


def places(location, limit=None, include_parens=False):
    """Tách cột Location thành danh sách địa danh."""
    if location is None or (not isinstance(location, str) and pd.isna(location)):
        return []
    s = str(location)
    s = re.sub(r"[()]", ",", s) if include_parens else re.sub(r"\(.*?\)", "", s)
    out = []
    for part in re.split(r"[,;]| and ", s):
        part = PLACE_NOISE.sub("", part).strip(" .-")
        if len(part) >= 3 and part.lower() not in [p.lower() for p in out]:
            out.append(part)
    return out[:limit] if limit else out


PROVINCE_WORDS = {"Jawa": "Java", "Sumatera": "Sumatra", "Nangroe Aceh Darussalam": "Aceh",
                  "Dki Jakarta": "Jakarta", "Kepulauan": "Islands"}
DIRECTIONS = {"Barat": "West", "Timur": "East", "Tengah": "Central", "Selatan": "South", "Utara": "North"}


def anglicize(name):
    """Tên tỉnh Indonesia -> cách báo tiếng Anh hay viết: 'Jawa Barat' -> 'West Java'."""
    for k, v in PROVINCE_WORDS.items():
        name = re.sub(rf"\b{k}\b", v, name, flags=re.I)
    words = name.split()
    if len(words) >= 2 and words[-1].capitalize() in DIRECTIONS:
        words = [DIRECTIONS[words[-1].capitalize()]] + words[:-1]
    return " ".join(words)


def provinces(location):
    """Lấy tên tỉnh dạng '... (X province)' trong cột Location (định dạng của Indonesia/Philippines)."""
    if location is None or (not isinstance(location, str) and pd.isna(location)):
        return []
    out = []
    for grp in re.findall(r"\(([^()]*province[^()]*)\)", str(location), flags=re.I):
        for part in re.split(r"[,;/]", PLACE_NOISE.sub("", grp)):
            part = part.strip(" .-")
            if len(part) >= 3 and part.lower() not in [x.lower() for x in out]:
                out.append(part)
    return out


GENERIC_WORDS = ["storm", "typhoon", "tropical", "depression", "cyclone", "super", "severe",
                 "flood", "floods", "the", "no", "number", "slide", "landslide"]
NOT_A_NAME = set(GENERIC_WORDS) | {"rain", "rains", "flooding", "heavy"}


def _is_generic(w):
    return w in GENERIC_WORDS or bool(difflib.get_close_matches(w, GENERIC_WORDS, n=1, cutoff=0.82))


def event_names(row):
    """Mọi tên bão (quốc tế + địa phương) trong cột Event Name.
    "Typhoon 'Damrey' / 'Ramil'" -> ['damrey','ramil'];  "Cyclone 'Seroja-21'" -> ['seroja-21','seroja'];
    Sự kiện tên chung chung như 'slide' -> [] (coi như không có tên)."""
    event = _cell(row, "Event Name")
    if not event:
        return []
    names = []
    for piece in re.split(r"[()/&,;]|\band\b", event.replace("\u2019", "'")):
        words = [w for w in re.findall(r"[a-z0-9-]+", _fold(piece).replace("'", " ")) if not _is_generic(w)]
        if not words:
            continue
        full = " ".join(words)
        for n in (full, re.sub(r"-\d+$", "", full)):
            if len(n) >= 3 and not n.isdigit() and n not in NOT_A_NAME and n not in names:
                names.append(n)
    return names


HAZARD_SYNONYMS = {
    "Flood": ["flood", "flooding"],
    "Storm": ["typhoon", "tropical storm"],
    "Mass movement (wet)": ["landslide"],
    "Drought": ["drought"],
    "Extreme temperature": ["heat wave", "cold spell"],
    "Wildfire": ["wildfire"],
}

HAZARD_RE = re.compile(
    r"\b(storms?|typhoons?|cyclones?|depressions?|floods?|flooding|flooded|inundat\w*|landslides?|"
    r"mudslides?|heavy rain\w*|rainfall|drought|heat ?waves?|cold (?:spell|wave|snap)|saltwater|"
    r"salinity|disasters?|damages?d?)\b", re.I)

TYPE_RE = {
    "Flood": re.compile(r"\b(floods?|flooding|flooded|inundat\w*|heavy rain\w*|rainfall|overflow\w*)\b", re.I),
    "Storm": re.compile(r"\b(storms?|typhoons?|cyclones?|depressions?|gales?)\b", re.I),
    "Mass movement (wet)": re.compile(r"\b(landslides?|mudslides?|mass movement|slope failure)\b", re.I),
    "Drought": re.compile(r"\b(drought|dry spell|water shortage|saltwater|salinity)\b", re.I),
    "Extreme temperature": re.compile(r"\b(heat ?waves?|cold (?:spell|wave|snap)|extreme (?:heat|cold)|temperatures?)\b", re.I),
}


# ======================= PHÂN BIỆT SỰ KIỆN: SUBTYPE / ORIGIN / ASSOCIATED TYPES =======================
# (khóa nhận diện trong cột EM-DAT, [từ khóa thường xuất hiện trong bài báo]).
# Khóa CỤ THỂ đặt trước khóa chung; mỗi cụm trong ô chỉ khớp khóa đầu tiên tìm thấy.
HAZARD_KW = [
    ("flash flood",     ["flash flood", "flash floods", "flash-flood", "flash-floods"]),
    ("riverine",        ["riverine", "river flood", "river floods", "river overflow", "river burst",
                         "overflowing river", "rivers overflowed", "burst its banks", "burst their banks"]),
    ("storm surge",     ["storm surge", "tidal surge", "coastal flood", "coastal flooding"]),
    ("coastal",         ["coastal flood", "coastal flooding", "storm surge", "tidal flood", "high tide",
                         "king tide", "sea level"]),
    ("flood",           ["flood", "floods", "flooding", "flooded", "inundated"]),
    ("tropical cyclone", ["typhoon", "tropical storm", "tropical cyclone", "tropical depression",
                          "cyclone", "super typhoon", "low pressure area"]),
    ("convective",      ["thunderstorm", "thunderstorms", "lightning", "tornado", "hailstorm", "squall",
                         "whirlwind", "waterspout", "windstorm", "strong winds", "gusts"]),
    ("extra-tropical",  ["extratropical", "extra-tropical", "northeast monsoon", "southwest monsoon",
                         "cold surge", "monsoon trough"]),
    ("extratropical",   ["extratropical", "extra-tropical", "northeast monsoon", "southwest monsoon",
                         "cold surge", "monsoon trough"]),
    ("storm",           ["storm", "storms", "gale", "gales"]),
    ("mudslide",        ["mudslide", "mudslides", "mudflow", "debris flow", "lahar", "mud flow"]),
    ("landslide",       ["landslide", "landslides", "landslip", "slope failure", "earth slip"]),
    ("mass movement",   ["landslide", "landslides", "mudslide", "slope failure"]),
    ("ground movement", ["landslide", "sinkhole", "ground subsidence"]),
    ("rockfall",        ["rockfall", "rock fall", "falling rocks"]),
    ("avalanche",       ["avalanche", "snow slide"]),
    ("heat",            ["heat wave", "heatwave", "extreme heat", "scorching", "record temperature",
                         "heat index"]),
    ("cold",            ["cold wave", "cold spell", "cold snap", "frost", "freezing", "cold front"]),
    ("winter",          ["severe winter", "cold wave", "cold spell", "frost", "snowfall"]),
    ("drought",         ["drought", "dry spell", "water shortage", "el nino", "saltwater intrusion",
                         "salinity", "crop failure"]),
    ("fire",            ["forest fire", "wildfire", "bushfire", "peatland fire", "haze", "hotspots"]),
    # ---- các giá trị hay gặp ở cột Origin ----
    ("heavy rain",      ["heavy rain", "heavy rains", "heavy rainfall", "torrential rain", "downpour",
                         "downpours", "incessant rain", "prolonged rain"]),
    ("monsoon",         ["monsoon", "monsoon rains", "habagat", "amihan"]),
    ("typhoon",         ["typhoon", "super typhoon"]),
    ("dam",             ["dam", "reservoir", "spillway", "dam release", "dam collapse", "dam break",
                         "released water"]),
    ("tide",            ["high tide", "tidal", "king tide"]),
    ("el nino",         ["el nino"]),
    ("la nina",         ["la nina"]),
    ("snowmelt",        ["snowmelt", "melting snow"]),
    ("rain",            ["rain", "rains", "rainfall", "downpour"]),
]

KW_STOP = {"general", "wet", "dry", "other", "unspecified", "unknown", "nan", "none", "na", "n/a",
           "snow", "debris", "rock", "and", "the", "of"}


def term_groups(text):
    """Cột EM-DAT (Subtype / Origin / Associated Types) -> list regex; mỗi regex = 1 cụm trong ô.
    VD 'Flash flood, Landslide' -> [regex(flash flood...), regex(landslide...)]."""
    if text is None or (not isinstance(text, str) and pd.isna(text)):
        return []
    groups, seen = [], set()
    for phrase in re.split(r"[,;/&()]|\band\b", _fold(text)):
        phrase = phrase.strip(" .-'\"")
        if len(phrase) < 3 or phrase in KW_STOP:
            continue
        terms = next((kws for key, kws in HAZARD_KW if key in phrase), None)
        if terms is None:
            if len(phrase.split()) > 4:
                continue
            terms = [phrase]
        pat = r"\b(?:" + "|".join(re.escape(_fold(t)) for t in terms) + r")\b"
        if pat not in seen:
            seen.add(pat)
            groups.append(re.compile(pat))
    return groups


def place_tokens(row):
    toks = set()
    for p in places(row.get("Location"), include_parens=True):
        for name in (p, anglicize(p)):
            t = _fold(name)
            if len(t) >= 4:
                toks.add(t)
    return toks


def build_profiles(df):
    """Hồ sơ cho MỌI sự kiện của nước (kể cả sự kiện ngoài --year/--limit) để so sánh với nhau.
    rivals[code] = các sự kiện cùng nước có cửa sổ ngày chồng lên nhau (kể cả chính nó)."""
    profiles = {}
    for _, row in df.iterrows():
        code = row["Disaster Name"]
        dtype = _cell(row, "Disaster Type")
        names = event_names(row)
        profiles[code] = {
            "code": code, "row": row,
            "window": disaster_window(row), "core": disaster_core(row),
            "names": names,
            "name_rx": re.compile(r"\b(?:" + "|".join(re.escape(n) for n in names) + r")\b") if names else None,
            "ptoks": place_tokens(row),
            "type_rx": TYPE_RE.get(dtype, HAZARD_RE),
            "sub": term_groups(_cell(row, "Disaster Subtype")),
            "origin": term_groups(_cell(row, "Origin")),
            "assoc": term_groups(_cell(row, "Associated Types", "Associated Disaster Types", "Associated Dis")),
            "rivals": [],
        }
    for code, p in profiles.items():
        w = p["window"]
        p["rivals"] = [code] + [
            c for c, q in profiles.items()
            if c != code and w and q["window"] and w[0] <= q["window"][1] and q["window"][0] <= w[1]]
    return profiles


def prep(content):
    """Tính sẵn bản 'bỏ dấu + lower' của bài (dùng cho mọi phép so khớp)."""
    if "_blob" not in content:
        content["_title"] = _fold(content["title"])
        content["_blob"] = _fold(f"{content['title']} {content['full_text']}")
    return content


def is_relevant(content, prof):
    """Nhận bài nếu thỏa MỘT trong hai nhánh:
    (a) thảm họa có tên bão: bài nhắc tên bão VÀ có từ thiên tai (HAZARD_RE);
    (b) bài nói đúng LOẠI thiên tai VÀ (nhắc một địa danh trong Location HOẶC tiêu đề nói về loại thiên tai đó).
    Nhánh (b) áp dụng cho mọi thảm họa, kể cả bão có tên, nên bài chỉ viết "the typhoon" vẫn được giữ.
    NAMED_EVENT_STRICT = True: bão có tên chỉ xét nhánh (a) (luật cũ)."""
    prep(content)
    blob = content["_blob"]
    name_ok = (bool(prof["name_rx"]) and bool(prof["name_rx"].search(blob))
               and bool(HAZARD_RE.search(blob)))
    if prof["name_rx"] and NAMED_EVENT_STRICT:
        return name_ok
    if name_ok:                                       # nhánh (a)
        return True
    rx = prof["type_rx"]                              # nhánh (b)
    if not rx.search(blob):
        return False
    if not STRICT_PLACE:
        return True
    toks = prof["ptoks"]
    if not toks or any(t in blob for t in toks):
        return True
    return bool(rx.search(content["_title"]))


def date_score(date_str, core):
    """4 điểm nếu bài nằm trong thời gian thảm họa, giảm dần về 0 khi cách xa 60 ngày."""
    if not core or not date_str:
        return 2.0
    try:
        d = datetime.strptime(str(date_str)[:10], "%Y-%m-%d")
    except ValueError:
        return 2.0
    gap = 0 if core[0] <= d <= core[1] else min(abs((d - core[0]).days), abs((d - core[1]).days))
    return 4.0 - 4.0 * min(gap, 60) / 60


def score_article(content, p, named_hits):
    """Điểm 'bài này thuộc sự kiện p đến mức nào'."""
    blob, title = content["_blob"], content["_title"]
    s = 0.0
    if p["name_rx"]:                                              # tên bão: tín hiệu mạnh nhất
        if p["name_rx"].search(title):
            s += 12
        elif p["name_rx"].search(blob):
            s += 10
    if p["type_rx"].search(blob):                                 # Disaster Type
        s += 2
    s += 3.0 * min(2, sum(1 for g in p["sub"] if g.search(blob)))      # Disaster Subtype
    s += 2.0 * min(2, sum(1 for g in p["origin"] if g.search(blob)))   # Origin
    s += 1.5 * min(3, sum(1 for g in p["assoc"] if g.search(blob)))    # Associated Types
    toks = p["ptoks"]
    if toks:                                                      # Location
        s += min(4, sum(1 for t in toks if t in blob))
        if any(t in title for t in toks):
            s += 1.5
    s += date_score(content["date"], p["core"])                   # Start/End date
    if named_hits and p["code"] not in named_hits:                # bài đang nói về bão có tên của sự kiện khác
        s -= OTHER_NAMED_PENALTY
    return s


def pick_event(content, code, profiles):
    """Bài này thuộc sự kiện nào trong nhóm cạnh tranh (cùng nước, cùng thời gian)?
    Trả về (code_thắng | None, thông tin mơ hồ | None). None = không sự kiện nào nhận bài."""
    prep(content)
    rivals = profiles[code]["rivals"]
    named_hits = {rc for rc in rivals
                  if profiles[rc]["name_rx"] and profiles[rc]["name_rx"].search(content["_blob"])}
    cands = []
    for rc in rivals:
        p = profiles[rc]
        if date_ok(content["date"], p["window"]) and is_relevant(content, p):
            cands.append((score_article(content, p, named_hits), rc))
    if not cands:
        return None, None
    cands.sort(key=lambda x: (-x[0], x[1]))
    amb = None
    if len(cands) > 1 and cands[0][0] - cands[1][0] < AMBIGUOUS_MARGIN:
        amb = {"assigned": cands[0][1], "score": round(cands[0][0], 1),
               "runner_up": cands[1][1], "runner_score": round(cands[1][0], 1)}
    return cands[0][1], amb


def log_ambiguous(country, outlet, content, url, amb):
    """Ghi các bài khó phân biệt (2 sự kiện điểm gần nhau) ra file riêng để kiểm tay."""
    path = f"ambiguous_{country}.csv"
    rec = {"URL": url, "OUTLET": outlet, "Tittle": content["title"], "Date": content["date"],
           "Assigned to": amb["assigned"], "Score": amb["score"],
           "Runner-up": amb["runner_up"], "Runner-up score": amb["runner_score"]}
    pd.DataFrame([rec]).to_csv(path, mode="a", header=not os.path.exists(path),
                               index=False, encoding="utf-8-sig")


# ======================= TẠO QUERY =======================
def build_queries(row, site, country_name):
    """Nhiều biến thể query cho 1 đài. Ưu tiên: tên bão -> Subtype/Origin/Associated -> tỉnh -> cả nước."""
    dtype = _cell(row, "Disaster Type")
    syns = HAZARD_SYNONYMS.get(dtype, [dtype.lower()])
    year = _num(row.get("Start Year")) or ""
    month = _num(row.get("Start Month"))
    month_name = datetime(2000, month, 1).strftime("%B") if month and 1 <= month <= 12 else ""
    loc = row.get("Location")
    provs = [anglicize(x) for x in provinces(loc)]
    pl = (provs or places(loc))[:MAX_PLACES]

    def first_phrase(*cols):
        txt = re.sub(r"\(.*?\)", "", _cell(row, *cols))
        for part in re.split(r"[,;/&]", txt):
            part = part.strip().lower()
            if len(part) >= 4 and part not in KW_STOP:
                return part
        return ""

    sub = first_phrase("Disaster Subtype")
    origin = first_phrase("Origin")
    assoc = first_phrase("Associated Types", "Associated Disaster Types", "Associated Dis")

    qs = []
    for n in event_names(row)[:3]:                      # tên quốc tế + tên địa phương
        qs.append(f'"{n}" {syns[0]} {year} site:{site}')
    for place in pl:
        qs.append(f"{syns[0]} {place} {month_name} {year} site:{site}")
    if sub and sub != syns[0]:
        for place in pl[:1] or [country_name]:
            qs.append(f"{sub} {place} {month_name} {year} site:{site}")
    qs.append(f"{syns[0]} {country_name} {month_name} {year} site:{site}")
    qs.append(f"{syns[0]}s {country_name} {year} {month_name} deaths displaced site:{site}")
    if origin:
        qs.append(f"{origin} {syns[0]} {pl[0] if pl else country_name} {month_name} {year} site:{site}")
    if assoc and assoc not in (sub, syns[0]):
        qs.append(f"{assoc} {pl[0] if pl else country_name} {month_name} {year} site:{site}")
    for syn in syns[1:]:
        qs.append(f"{syn} {country_name} {month_name} {year} site:{site}")
    if pl:
        qs.append(f"{syns[0]} {pl[0]} {year} damage evacuated site:{site}")

    seen, uniq = set(), []
    for q in qs:
        q = re.sub(r"\s+", " ", q).strip()
        if q not in seen:
            seen.add(q)
            uniq.append(q)
    return uniq[:MAX_QUERIES_PER_OUTLET] if MAX_QUERIES_PER_OUTLET else uniq


# ======================= QUY TRÌNH CHÍNH =======================
def parse_years(text):
    if not text:
        return None
    m = re.fullmatch(r"\s*(\d{4})\s*(?:-\s*(\d{4}))?\s*", text)
    if not m:
        sys.exit("❌ --year phải có dạng 2016 hoặc 2016-2018")
    lo = int(m.group(1))
    return lo, int(m.group(2) or lo)


def load_country_disasters(df_all, code, years, limit):
    df = df_all.dropna(subset=["Disaster Name"]).copy()
    df["Disaster Name"] = df["Disaster Name"].astype(str).str.strip()
    df = df[df["Disaster Name"].str.endswith("-" + code)]
    if years:
        sy = pd.to_numeric(df["Start Year"], errors="coerce")
        df = df[(sy >= years[0]) & (sy <= years[1])]
    return df.head(limit) if limit else df


def process_country(country, df_all, outlets, years, limit):
    info = COUNTRY_INFO[country]
    ckpt = f"checkpoint_{country}.csv"
    combos_file = f"checkpoint_{country}_done.txt"

    profiles = build_profiles(load_country_disasters(df_all, info["code"], None, None))   # TẤT CẢ sự kiện của nước
    df = load_country_disasters(df_all, info["code"], years, limit)                       # sự kiện sẽ chạy
    print(f"\n📂 {country}: {len(df)} thảm họa | cào HẾT {len(outlets)} đài, không giới hạn số bài")
    print("   Thứ tự đài: " + " → ".join(n for n, _ in outlets))

    # ---- nạp checkpoint (chạy lại sẽ tiếp tục, không cào lại) ----
    rows = (pd.read_csv(ckpt, dtype=str, keep_default_na=False).to_dict("records")
            if os.path.exists(ckpt) else [])
    seen_urls = {norm_url(r["URL"]) for r in rows if r.get("URL")}
    seen_titles = {_fold(r["Tittle"]) for r in rows if r.get("Tittle") and r["Tittle"] != "No Title"}
    done_combos = {(r["Disaster Name"], r["OUTLET"]) for r in rows}
    if os.path.exists(combos_file):
        with open(combos_file, encoding="utf-8") as f:
            done_combos |= {tuple(l.rstrip("\n").split("\t")) for l in f if "\t" in l}

    if rows or done_combos:
        n_ev = len({r["Disaster Name"] for r in rows})
        print(f"   ♻️ Tiếp tục từ checkpoint: đã có {len(rows)} bài của {n_ev} thảm họa, "
              f"{len(done_combos)} cặp (thảm họa-đài) đã xong -> sẽ bỏ qua phần này.")

    def count(code):
        return sum(1 for r in rows if r["Disaster Name"] == code)

    def crawl_outlet(row, code, window, outlet, site):
        """Cào HẾT một đài. Mỗi bài tìm được sẽ được gán cho sự kiện PHÙ HỢP NHẤT trong nhóm
        sự kiện cùng nước/cùng thời gian (không nhất thiết là sự kiện đang chạy)."""
        if (code, outlet) in done_combos:
            return
        tried, st, net_fail = set(), Counter(), False

        def consider(url, c):
            if site in ENGLISH_ONLY_SITES and not is_english(c["full_text"], MIN_EN_RATIO):
                st["không phải tiếng Anh"] += 1             # đài chỉ nhận bài tiếng Anh
                return
            tkey = _fold(c["title"])
            if c["title"] != "No Title" and tkey in seen_titles:
                st["đã có/trùng"] += 1                        # bài đăng lại
                return
            best, amb = pick_event(c, code, profiles)
            if best is None:
                st["ngoài khoảng ngày" if not date_ok(c["date"], window) else "không liên quan"] += 1
                return
            seen_urls.add(norm_url(url))
            seen_titles.add(tkey)
            rows.append({"OUTLET": outlet, "Disaster Name": best, "Article ID": "",
                         "Tittle": c["title"], "Full text": c["full_text"],
                         "Date": c["date"], "URL": url})
            if amb:
                log_ambiguous(country, outlet, c, url, amb)
            if best == code:
                st["LẤY ĐƯỢC"] += 1
                print(f"      ✅ {c['title'][:70]}" + ("  ⚠️ khó phân biệt" if amb else ""))
            else:
                st["chuyển sang sự kiện khác"] += 1
                print(f"      ↪️ [{best}] {c['title'][:60]}")

        # ---- Nguồn 1: API WordPress (nếu đài dùng WordPress) ----
        if USE_WP_API and wp_available(site):
            dtype = _cell(row, "Disaster Type")
            terms = event_names(row)[:3] or HAZARD_SYNONYMS.get(dtype, [dtype.lower()])
            for term in terms:
                print(f"   🧩 [{outlet}] WordPress API: '{term}' trong khoảng ngày của thảm họa")
                for item in wp_search(site, term, window):
                    key = norm_url(item["url"])
                    st["link tìm thấy"] += 1
                    if key in seen_urls or key in tried:
                        st["đã có/trùng"] += 1
                        continue
                    tried.add(key)
                    c, why = wp_to_content(item)
                    if c:
                        consider(item["url"], c)
                    else:
                        st[why] += 1

        # ---- Nguồn 2: công cụ tìm kiếm (site:) ----
        for query in build_queries(row, site, info["name"]):
            print(f"   🔍 [{outlet}] {query}")
            urls = search_article_urls(query, site)
            if urls is None:                          # lỗi mạng: query này chưa được tìm
                net_fail = True
                continue
            st["link tìm thấy"] += len(urls)
            todo = []
            for url in urls:
                key = norm_url(url)
                if key in seen_urls or key in tried:
                    st["đã có/trùng"] += 1
                    continue
                tried.add(key)
                todo.append(url)
            if not todo:
                continue
            with ThreadPoolExecutor(max_workers=DOWNLOAD_WORKERS) as pool:   # tải song song
                fetched = list(pool.map(extract_article_content, todo))
            for url, (c, why) in zip(todo, fetched):
                if c:
                    consider(url, c)
                else:
                    st[why] += 1
        got = st["LẤY ĐƯỢC"]
        print(f"   📈 [{outlet}] " + (" | ".join(f"{k}: {v}" for k, v in st.items()) or "không có link nào"))
        # lưu bài trước, đánh dấu "xong đài" sau -> mất điện giữa chừng vẫn an toàn
        pd.DataFrame(rows, columns=COLS).to_csv(ckpt, index=False)
        if net_fail:                                  # có query bị lỗi mạng -> KHÔNG đánh dấu xong
            print(f"   → {outlet}: +{got} bài (⚠️ có query lỗi mạng, đài này sẽ được chạy lại ở lần sau)")
            return
        done_combos.add((code, outlet))
        with open(combos_file, "a", encoding="utf-8") as f:
            f.write(f"{code}\t{outlet}\n")
        print(f"   → {outlet}: +{got} bài")

    report = []

    def _run_event(row):
        code = row["Disaster Name"]
        window = disaster_window(row)
        n_rivals = len(profiles[code]["rivals"]) - 1
        print(f"\n📌 [{code}] {row.get('Disaster Type')} | {row.get('Disaster Subtype')} | "
              f"{row.get('Event Name')} | {row.get('Location')}"
              + (f"  (cạnh tranh với {n_rivals} sự kiện cùng kỳ)" if n_rivals else ""))

        for outlet, site in outlets:                       # cào HẾT mọi đài, không dừng sớm
            crawl_outlet(row, code, window, outlet, site)

        n = count(code)
        by_outlet = "; ".join(f"{o}: {sum(1 for r in rows if r['Disaster Name'] == code and r['OUTLET'] == o)}"
                              for o, _ in outlets
                              if any(r["Disaster Name"] == code and r["OUTLET"] == o for r in rows))
        print(f"   📊 {code}: {n} bài ({by_outlet})")
        report.append({"Disaster Name": code, "Type": row.get("Disaster Type"),
                       "Subtype": row.get("Disaster Subtype"), "Event Name": row.get("Event Name"),
                       "Articles": n, "Status": "KHÔNG CÓ BÀI" if n == 0 else "OK",
                       "By outlet": by_outlet})

    for _, row in df.iterrows():
        try:
            _run_event(row)
        except KeyboardInterrupt:
            print("\n⏹️  Đã dừng theo yêu cầu. Đang lưu phần đã cào... (chạy lại đúng lệnh này để làm tiếp)")
            pd.DataFrame(rows, columns=COLS).to_csv(ckpt, index=False)
            break

    pd.DataFrame(report).to_csv(f"report_{country}.csv", index=False, encoding="utf-8-sig")

    # ---- sắp xếp theo thứ tự thảm họa -> thứ tự đài, rồi đánh Article ID ----
    ev_order = {c: i for i, c in enumerate(df["Disaster Name"])}
    ou_order = {o: i for i, (o, _) in enumerate(outlets)}
    out = pd.DataFrame(rows, columns=COLS)
    out = out[out["Disaster Name"].isin(ev_order)]         # chỉ giữ thảm họa đang chạy (--year/--limit)
    out = (out.assign(_e=out["Disaster Name"].map(ev_order), _o=out["OUTLET"].map(ou_order))
              .sort_values(["_e", "_o"], kind="stable").drop(columns=["_e", "_o"]).reset_index(drop=True))
    if AUTO_ARTICLE_ID:
        out["Article ID"] = [f"{info['prefix']}{i:03d}" for i in range(1, len(out) + 1)]
    else:
        out["Article ID"] = ""                            # để trống, đánh tay sau
    return out


def reassign_country(country, df_all):
    """Chấm lại TOÀN BỘ bài đã cào trong checkpoint bằng logic phân biệt sự kiện mới
    (Subtype/Origin/Associated Types/tên bão/địa danh/ngày) và chuyển bài về đúng sự kiện.
    Không xóa bài nào. Sau đó chạy lại lệnh cào bình thường để xuất lại file Excel."""
    info = COUNTRY_INFO[country]
    ckpt = f"checkpoint_{country}.csv"
    if not os.path.exists(ckpt):
        print(f"⏭️  {country}: chưa có {ckpt}")
        return
    profiles = build_profiles(load_country_disasters(df_all, info["code"], None, None))
    rows = pd.read_csv(ckpt, dtype=str, keep_default_na=False)
    shutil.copyfile(ckpt, ckpt + ".bak")
    moved = flagged = 0
    for i, r in rows.iterrows():
        code = r["Disaster Name"]
        if code not in profiles:
            continue
        c = {"title": r["Tittle"], "full_text": r["Full text"], "date": r["Date"]}
        best, amb = pick_event(c, code, profiles)
        if best and best != code:
            rows.at[i, "Disaster Name"] = best
            moved += 1
            print(f"   ↪️ {code} -> {best}: {r['Tittle'][:70]}")
        if amb:
            flagged += 1
            log_ambiguous(country, r["OUTLET"], c, r["URL"], amb)
    rows.to_csv(ckpt, index=False)
    print(f"✅ {country}: đã chuyển {moved} bài về đúng sự kiện, {flagged} bài khó phân biệt "
          f"(xem ambiguous_{country}.csv). Bản sao lưu: {ckpt}.bak")


def write_into_template(wb, country, df_out):
    """Ghi vào sheet đúng form: A=OUTLET, B=Disaster Name, C=Article ID, D=Tittle, E=Full text, F=Date, G=URL."""
    if country not in wb.sheetnames:
        print(f"⚠️ Sheet '{country}' không tồn tại trong file mẫu!")
        return
    ws = wb[country]
    if df_out.empty:
        print(f"⚠️ {country}: không có bài nào -> giữ nguyên sheet mẫu.")
        return
    first = find_header_row(ws) + 1
    styles = [copy(ws.cell(row=first, column=c)._style) for c in range(1, len(COLS) + 1)]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first:
            ws.unmerge_cells(str(rng))
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)
    for i, rec in enumerate(df_out[COLS].itertuples(index=False), start=first):
        for j, val in enumerate(rec, start=1):
            cell = ws.cell(row=i, column=j, value=ILLEGAL_CHARACTERS_RE.sub("", str(val)) if val != "" else None)
            cell._style = copy(styles[j - 1])              # giữ font/màu của form mẫu
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def run_country(country, years, limit):
    if not os.path.exists(DISASTER_FILE) or not os.path.exists(NEWS_TEMPLATE):
        sys.exit(f"❌ Thiếu '{DISASTER_FILE}' hoặc '{NEWS_TEMPLATE}' trong thư mục chạy.")
    df_all = pd.read_excel(DISASTER_FILE, sheet_name=DISASTER_SHEET)
    wb = load_workbook(NEWS_TEMPLATE)
    outlets = get_outlets(country)
    write_into_template(wb, country, process_country(country, df_all, outlets, years, limit))
    out_path = OUTPUT_PATTERN.format(country=country)
    try:
        wb.save(out_path)
    except PermissionError:                               # file đang mở trong Excel -> lưu sang tên khác
        out_path = out_path.replace(".xlsx", f" ({datetime.now():%H%M%S}).xlsx")
        wb.save(out_path)
        print(f"⚠️ File output đang mở trong Excel nên đã lưu sang '{out_path}'.")
    print(f"\n🎉 {country} xong! -> '{out_path}' (xem report_{country}.csv và ambiguous_{country}.csv)")


def merge_outputs():
    wb = load_workbook(NEWS_TEMPLATE)
    for country in COUNTRY_INFO:
        path = OUTPUT_PATTERN.format(country=country)
        if not os.path.exists(path):
            print(f"⏭️  Chưa có {path}")
            continue
        hdr = find_header_row(load_workbook(path, read_only=True)[country])
        df = pd.read_excel(path, sheet_name=country, header=hdr - 1, usecols=range(len(COLS)),
                           dtype=str, keep_default_na=False)
        df.columns = COLS
        df = df[df["URL"].str.strip() != ""]
        write_into_template(wb, country, df)
        print(f"✅ {country}: {len(df)} bài")
    wb.save(MERGED_OUTPUT)
    print(f"\n🎉 Đã gộp -> '{MERGED_OUTPUT}'")


def check_country(country):
    """Kiểm tra nhanh từng đài của 1 nước: tìm được link không, tải được bài không, có API WordPress không."""
    outlets = get_outlets(country)
    name = COUNTRY_INFO[country]["name"]
    print(f"\n🩺 KIỂM TRA {country} (mỗi đài ~10-30 giây)\n")
    print(f"{'ĐÀI':38} {'LINK':>5}  {'TẢI BÀI':14} WP-API")
    print("-" * 72)
    for outlet, site in outlets:
        urls = search_article_urls(f"flood {name} site:{site}", site)
        if urls is None:
            link, dl = "LỖI", "-"
        else:
            link, dl = str(len(urls)), "-"
            for u in urls[:2]:                        # thử tải tối đa 2 bài
                c, why = extract_article_content(u)
                if c:
                    dl = f"OK ({len(c['full_text'].split())} từ)"
                    break
                dl = why
        wp = "có" if (USE_WP_API and wp_available(site)) else "không"
        print(f"{outlet[:37]:38} {link:>5}  {dl:14} {wp}")
    print("\nĐọc kết quả: LINK=0 -> tìm kiếm không ra (đổi SEARCH_BACKENDS / kiểm tra mạng); 'tải lỗi' -> báo chặn bot;"
          "\n'quá ngắn' -> paywall; WP-API 'có' -> script lấy bài qua API rất nhanh.")


def main():
    ap = argparse.ArgumentParser(description="Cào tin thiên tai vào form NEWS DATA (mỗi người 1 nước).")
    ap.add_argument("countries", nargs="*",
                    help="INDONESIA | PHILIPPINES | THAILAND | MALAYSIA | VIETNAM")
    ap.add_argument("--year", help="Lọc năm bắt đầu thảm họa: 2016 hoặc 2016-2018 (mặc định: tất cả)")
    ap.add_argument("--limit", type=int, help="Chỉ chạy N thảm họa đầu (chạy thử)")
    ap.add_argument("--merge", action="store_true", help="Gộp file các nước thành 1 file")
    ap.add_argument("--check", action="store_true", help="Kiểm tra nhanh các đài của nước đó (không cào thật)")
    ap.add_argument("--reassign", action="store_true",
                    help="Chấm lại các bài đã cào trong checkpoint và chuyển về đúng sự kiện (không cào mới)")
    args = ap.parse_args()

    if args.merge:
        return merge_outputs()
    if not args.countries:
        ap.print_help()
        sys.exit(1)
    for c in (x.upper() for x in args.countries):
        if c not in COUNTRY_INFO:
            sys.exit(f"❌ Nước '{c}' không hợp lệ. Chọn: {', '.join(COUNTRY_INFO)}")
        if args.check:
            check_country(c)
            continue
        if args.reassign:
            reassign_country(c, pd.read_excel(DISASTER_FILE, sheet_name=DISASTER_SHEET))
            continue
        run_country(c, parse_years(args.year), args.limit)


if __name__ == "__main__":
    main()