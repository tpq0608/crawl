# -*- coding: utf-8 -*-
"""Các lệnh chạy từ dòng lệnh: cào 1 nước, chấm lại checkpoint, gộp file, kiểm tra đài."""
import os
import shutil
import sys
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from .config import (DISASTER_FILE, DISASTER_SHEET, NEWS_TEMPLATE, OUTPUT_PATTERN, MERGED_OUTPUT,
                     COUNTRY_INFO, COLS, USE_WP_API)
from .search import search_article_urls
from .fetch import extract_article_content, wp_available
from .matching import build_profiles, pick_event, log_ambiguous
from .excel_io import find_header_row, write_into_template
from .pipeline import get_outlets, load_country_disasters, process_country


def reassign_country(country, df_all):
    """Chấm lại TOÀN BỘ bài đã cào trong checkpoint bằng logic phân biệt sự kiện mới
    (Subtype/Origin/Associated Types/tên bão/địa danh/ngày) và chuyển bài về đúng sự kiện.
    Không xóa bài nào. Sau đó chạy lại lệnh cào bình thường để xuất lại file Excel."""
    info = COUNTRY_INFO[country]
    ckpt = f"checkpoint_{country}.csv"
    if not os.path.exists(ckpt):
        print(f"⏭️  {country}: chưa có {ckpt}")
        return
    profiles = build_profiles(load_country_disasters(df_all, info["code"], None, None))
    rows = pd.read_csv(ckpt, dtype=str, keep_default_na=False)
    shutil.copyfile(ckpt, ckpt + ".bak")
    moved = flagged = 0
    for i, r in rows.iterrows():
        code = r["Disaster Name"]
        if code not in profiles:
            continue
        c = {"title": r["Tittle"], "full_text": r["Full text"], "date": r["Date"]}
        best, amb = pick_event(c, code, profiles)
        if best and best != code:
            rows.at[i, "Disaster Name"] = best
            moved += 1
            print(f"   ↪️ {code} -> {best}: {r['Tittle'][:70]}")
        if amb:
            flagged += 1
            log_ambiguous(country, r["OUTLET"], c, r["URL"], amb)
    rows.to_csv(ckpt, index=False)
    print(f"✅ {country}: đã chuyển {moved} bài về đúng sự kiện, {flagged} bài khó phân biệt "
          f"(xem ambiguous_{country}.csv). Bản sao lưu: {ckpt}.bak")


def run_country(country, years, limit):
    if not os.path.exists(DISASTER_FILE) or not os.path.exists(NEWS_TEMPLATE):
        sys.exit(f"❌ Thiếu '{DISASTER_FILE}' hoặc '{NEWS_TEMPLATE}' trong thư mục chạy.")
    df_all = pd.read_excel(DISASTER_FILE, sheet_name=DISASTER_SHEET)
    wb = load_workbook(NEWS_TEMPLATE)
    outlets = get_outlets(country)
    write_into_template(wb, country, process_country(country, df_all, outlets, years, limit))
    out_path = OUTPUT_PATTERN.format(country=country)
    try:
        wb.save(out_path)
    except PermissionError:                               # file đang mở trong Excel -> lưu sang tên khác
        out_path = out_path.replace(".xlsx", f" ({datetime.now():%H%M%S}).xlsx")
        wb.save(out_path)
        print(f"⚠️ File output đang mở trong Excel nên đã lưu sang '{out_path}'.")
    print(f"\n🎉 {country} xong! -> '{out_path}' (xem report_{country}.csv và ambiguous_{country}.csv)")


def merge_outputs():
    wb = load_workbook(NEWS_TEMPLATE)
    for country in COUNTRY_INFO:
        path = OUTPUT_PATTERN.format(country=country)
        if not os.path.exists(path):
            print(f"⏭️  Chưa có {path}")
            continue
        hdr = find_header_row(load_workbook(path, read_only=True)[country])
        df = pd.read_excel(path, sheet_name=country, header=hdr - 1, usecols=range(len(COLS)),
                           dtype=str, keep_default_na=False)
        df.columns = COLS
        df = df[df["URL"].str.strip() != ""]
        write_into_template(wb, country, df)
        print(f"✅ {country}: {len(df)} bài")
    wb.save(MERGED_OUTPUT)
    print(f"\n🎉 Đã gộp -> '{MERGED_OUTPUT}'")


def check_country(country):
    """Kiểm tra nhanh từng đài của 1 nước: tìm được link không, tải được bài không, có API WordPress không."""
    outlets = get_outlets(country)
    name = COUNTRY_INFO[country]["name"]
    print(f"\n🩺 KIỂM TRA {country} (mỗi đài ~10-30 giây)\n")
    print(f"{'ĐÀI':38} {'LINK':>5}  {'TẢI BÀI':14} WP-API")
    print("-" * 72)
    for outlet, site in outlets:
        urls = search_article_urls(f"flood {name} site:{site}", site)
        if urls is None:
            link, dl = "LỖI", "-"
        else:
            link, dl = str(len(urls)), "-"
            for u in urls[:2]:                        # thử tải tối đa 2 bài
                c, why = extract_article_content(u)
                if c:
                    dl = f"OK ({len(c['full_text'].split())} từ)"
                    break
                dl = why
        wp = "có" if (USE_WP_API and wp_available(site)) else "không"
        print(f"{outlet[:37]:38} {link:>5}  {dl:14} {wp}")
    print("\nĐọc kết quả: LINK=0 -> tìm kiếm không ra (đổi SEARCH_BACKENDS / kiểm tra mạng); 'tải lỗi' -> báo chặn bot;"
          "\n'quá ngắn' -> paywall; WP-API 'có' -> script lấy bài qua API rất nhanh.")
