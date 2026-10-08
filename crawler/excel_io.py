# -*- coding: utf-8 -*-
"""Đọc/ghi file Excel theo form mẫu NEWS DATA."""
from copy import copy

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment

from .config import FIRST_DATA_ROW, COLS


def find_header_row(ws):
    """Dòng tiêu đề = dòng có ô cột A là 'OUTLET' (mặc định dòng 2). Dữ liệu bắt đầu ngay dưới nó."""
    for r in range(1, 11):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().upper() == "OUTLET":
            return r
    return FIRST_DATA_ROW - 1


def write_into_template(wb, country, df_out):
    """Ghi vào sheet đúng form: A=OUTLET, B=Disaster Name, C=Article ID, D=Tittle, E=Full text, F=Date, G=URL."""
    if country not in wb.sheetnames:
        print(f"⚠️ Sheet '{country}' không tồn tại trong file mẫu!")
        return
    ws = wb[country]
    if df_out.empty:
        print(f"⚠️ {country}: không có bài nào -> giữ nguyên sheet mẫu.")
        return
    first = find_header_row(ws) + 1
    styles = [copy(ws.cell(row=first, column=c)._style) for c in range(1, len(COLS) + 1)]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first:
            ws.unmerge_cells(str(rng))
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)
    for i, rec in enumerate(df_out[COLS].itertuples(index=False), start=first):
        for j, val in enumerate(rec, start=1):
            cell = ws.cell(row=i, column=j, value=ILLEGAL_CHARACTERS_RE.sub("", str(val)) if val != "" else None)
            cell._style = copy(styles[j - 1])              # giữ font/màu của form mẫu
            cell.alignment = Alignment(vertical="top", wrap_text=False)
